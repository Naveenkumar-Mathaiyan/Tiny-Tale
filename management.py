"""Additive staff access, stock-only editing and date-filtered exports."""
import io, re, uuid, secrets
from datetime import datetime, timedelta, timezone
from collections import Counter, defaultdict
from html import escape
from flask import request, session, jsonify, send_file
from sqlalchemy import select, text
from sqlalchemy.exc import IntegrityError
from werkzeug.security import generate_password_hash

IST=timedelta(hours=5,minutes=30)
def install(m):
 app,db=m.app,m.db
 class Staff(db.Model):
  id=db.Column(db.String(36),primary_key=True)
  username=db.Column(db.String(60),unique=True,nullable=False)
  name=db.Column(db.String(100),nullable=False)
  role=db.Column(db.String(20),nullable=False)
  password_hash=db.Column(db.Text,nullable=False)
  active=db.Column(db.Boolean,default=True,nullable=False)
  version=db.Column(db.Integer,default=1,nullable=False)
  created=db.Column(db.DateTime,default=datetime.utcnow)
 with app.app_context():
  guard=db.engine.connect() if db.engine.dialect.name=='postgresql' else None
  try:
   if guard:guard.execute(text('SELECT pg_advisory_lock(846304)'))
   db.create_all()
  finally:
   if guard:guard.execute(text('SELECT pg_advisory_unlock(846304)'));guard.close()
 m.Staff=Staff
 def identity():
  if not session.get('admin'):return None
  if not session.get('staff_id'):return dict(name='Owner',username='',role='super_admin')
  s=db.session.get(Staff,session['staff_id'])
  if not s or not s.active or s.version!=session.get('staff_version'):return None
  return dict(id=s.id,name=s.name,username=s.username,role=s.role)
 m.staff_identity=identity
 @app.before_request
 def access_control():
  if not request.path.startswith('/api/admin/') or request.path=='/api/admin/login':return
  if m.ROLE!='admin':return jsonify(error='Not found'),404
  user=identity()
  if not user:session.clear();return jsonify(error='Please sign in again.'),401
  endpoint=request.path.removeprefix('/api/admin/').split('/')[0]
  if endpoint=='accounts' and user['role']!='super_admin':return jsonify(error='Only the super admin can manage staff accounts.'),403
  if user['role']=='store_keeper':
   allowed={'session','logout','stock','upload','assets','merch'}
   if endpoint=='products' and request.method=='GET':return
   if endpoint not in allowed:return jsonify(error='Your account can update stock and product media only.'),403
 @app.route('/api/admin/accounts',methods=['GET','POST'])
 @m.admin
 def accounts():
  if request.method=='GET':return jsonify([dict(id=s.id,username=s.username,name=s.name,role=s.role,active=s.active) for s in db.session.scalars(select(Staff).order_by(Staff.name))])
  d=request.get_json(silent=True) or {};sid=d.get('id');s=db.session.get(Staff,sid) if sid else None
  if sid and not s:return jsonify(error='Account not found.'),404
  try:
   username=str(d.get('username','')).strip().lower();name=str(d.get('name','')).strip();role=d.get('role');password=d.get('password','');active=d.get('active',True)
   if not re.fullmatch(r'[a-z0-9][a-z0-9._-]{2,59}',username):raise ValueError('Username: 3–60 letters, digits, dots, underscores or hyphens.')
   if not name or len(name)>100 or role not in ['admin','store_keeper'] or type(active)is not bool:raise ValueError('Enter a name, valid role and active status.')
   if not isinstance(password,str) or (password and not 12<=len(password)<=128) or (not s and not password):raise ValueError('New passwords must contain 12–128 characters.')
   existing=db.session.scalar(select(Staff).where(Staff.username==username))
   if existing and existing.id!=sid:raise ValueError('This username is already used.')
   if not s:s=Staff(id=str(uuid.uuid4()),version=0);db.session.add(s)
   s.username=username;s.name=name;s.role=role;s.active=active;s.version+=1
   if password:s.password_hash=generate_password_hash(password)
   db.session.commit();return jsonify(ok=True)
  except ValueError as e:db.session.rollback();return jsonify(error=str(e)),400
  except IntegrityError:db.session.rollback();return jsonify(error='This username is already used.'),400
 @app.post('/api/admin/stock')
 @m.admin
 def stock():
  d=request.get_json(silent=True) or {};p=db.session.scalar(select(m.Product).where(m.Product.id==d.get('id')).with_for_update())
  if not p:return jsonify(error='Product not found.'),404
  try:
   entries=d.get('sizes');current=list(db.session.scalars(select(m.ProductSize).where(m.ProductSize.product_id==p.id)))
   if not isinstance(entries,list) or any(not isinstance(x,dict) for x in entries) or len(entries)!=len(current) or {x.get('size') for x in entries}!={x.size for x in current}:raise ValueError('Keep the existing size configuration; ask an admin to add or remove sizes.')
   for x in entries:
    row=next(v for v in current if v.size==x['size'])
    if type(x.get('stock'))is not int or not 0<=x['stock']<=100000 or x.get('enabled')!=row.enabled:raise ValueError('Use stock from 0–100000; store keepers cannot change enabled sizes.')
   image=d.get('image',p.image)
   if image!=p.image:
    if not isinstance(image,str) or not image.startswith('/media/') or not db.session.get(m.Media,image.removeprefix('/media/')):raise ValueError('Upload a product cover image first.')
   # Reject extra properties rather than silently accepting attempted price/content edits.
   if set(d)-{'id','sizes','image'}:return jsonify(error='Only stock and cover images may be changed here.'),403
   for x in entries:next(v for v in current if v.size==x['size']).stock=x['stock']
   p.image=image;m.sync_stock(p);db.session.commit();return jsonify(p.data(private=True))
  except (ValueError,TypeError,KeyError,StopIteration) as e:db.session.rollback();return jsonify(error=str(e) or 'Invalid stock values.'),400

 def dates():
  try:
   start=datetime.strptime(request.args.get('from',''),'%Y-%m-%d');end=datetime.strptime(request.args.get('to',''),'%Y-%m-%d')
   if end<start or (end-start).days>365:raise ValueError()
   return start-IST,end+timedelta(days=1)-IST,start.strftime('%Y-%m-%d'),end.strftime('%Y-%m-%d')
  except ValueError:raise ValueError('Choose valid From/To dates, up to 366 inclusive days.')
 def rows_for(model,start,end):
  rows=list(db.session.scalars(select(model).where(model.created>=start,model.created<end).order_by(model.created).limit(5001)))
  if len(rows)>5000:raise ValueError('More than 5000 rows. Choose a shorter date range.')
  return rows
 def stamp(dt):return (dt.replace(tzinfo=None)+IST).strftime('%d %b %Y %H:%M:%S')
 def report_data(kind,start,end):
  note='Dates and times use IST. Test payments collect no money.'
  if kind=='orders':
   headers=['Order number','Date (IST)','Customer','Status','Payment mode','Method','Total (INR)'];data=[];confirmed=0;tests=0
   for o in rows_for(m.Order,start,end):
    meta=db.session.get(m.PurchaseMeta,o.id);mode=meta.mode if meta else 'legacy';tests+=mode=='test'
    if mode!='test' and o.status in ['Confirmed','Shipped','Delivered']:confirmed+=o.total
    data.append([meta.number if meta else o.id,stamp(o.created),o.customer.get('name',''),o.status,mode,meta.payment_method if meta else '',o.total])
   note+=f' {len(data)} orders/enquiries; {tests} test orders. Confirmed non-test order value: INR {confirmed}. This is not a payment-settlement report.'
  elif kind=='inventory':
   headers=['Product code','Product','Size','Stock','Visibility','Stock status'];data=[]
   for p,v in db.session.execute(select(m.Product,m.ProductSize).join(m.ProductSize).where(m.ProductSize.enabled==True).order_by(m.Product.name,m.ProductSize.position).limit(5001)):
    data.append([p.id,p.name,v.size,v.stock,'Active' if p.active else 'Hidden','Restock soon' if v.stock==0 else 'Low stock' if v.stock<5 else 'Available'])
   note='Current inventory snapshot at generation time. Date range does not reconstruct past stock.'
  elif kind=='restock':
   headers=['Requested (IST)','Product','Size','Customer email','Quantity','Status'];data=[]
   for r in rows_for(m.Restock,start,end):
    p=db.session.get(m.Product,r.product_id);c=db.session.get(m.Customer,r.customer_id);data.append([stamp(r.created),p.name if p else r.product_id,r.size,c.email if c else '',r.qty,r.status])
   note+=' Repeat requests may update an existing request; date is its original creation date.'
  elif kind=='traffic':
   headers=['Date (IST)','Source','Medium','Campaign','Country','State','Landing page'];data=[]
   for r in rows_for(m.Traffic,start,end):data.append([stamp(r.created),r.source,r.medium,r.campaign or 'No campaign tag',r.country,r.state,r.landing])
   note='Consented first-touch browser sessions; these are not unique customers. Traffic is retained for 90 days; older ranges may be incomplete.'
  elif kind=='journey':
   headers=['Session started (IST)','Session reference','Navigation','Active seconds','Last stage'];data=[];groups={}
   events=rows_for(m.Journey,start,end)
   for e in events:
    g=groups.setdefault(e.sid,dict(start=e.created,path=[],seconds=0,last=e.stage))
    if e.kind!='active' and (e.stage!='success' or e.kind=='test_payment_success') and (not g['path'] or g['path'][-1]!=e.stage):g['path'].append(e.stage)
    g['seconds']+=e.seconds;g['last']=e.stage
   for sid,g in groups.items():data.append([stamp(g['start']),sid[:8],' > '.join(g['path']),g['seconds'],g['last']])
   note='Consented activity within the chosen dates. Session start is the first recorded event in this range. Active time excludes idle/background time.'
  else:raise ValueError('Choose orders, inventory, restock, traffic or journey.')
  if len(data)>5000:raise ValueError('More than 5000 rows. Choose a smaller report.')
  return headers,data,note
 m.report_data=report_data
 @app.get('/api/admin/reports')
 @m.admin
 def reports():
  try:
   start,end,frm,to=dates();kind=request.args.get('type','orders');fmt=request.args.get('format','pdf')
   if fmt not in ['pdf','xlsx']:raise ValueError('Choose PDF or Excel.')
   headers,data,note=report_data(kind,start,end)
  except ValueError as e:return jsonify(error=str(e)),400
  filename=f'TinyTale-{kind}-{frm}-{to}.{fmt}';buf=io.BytesIO()
  if fmt=='xlsx':
   from openpyxl import Workbook
   from openpyxl.styles import Font, PatternFill
   from openpyxl.utils import get_column_letter
   wb=Workbook();ws=wb.active;ws.title='Report';ws.append(headers)
   for row in data:
    ws.append(row)
    for cell in ws[ws.max_row]:
     if isinstance(cell.value,str):cell.data_type='s' # Never interpret uploaded/user values as Excel formulas.
   for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='123E52')
   ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
   for i,h in enumerate(headers,1):ws.column_dimensions[get_column_letter(i)].width=min(42,max(16,len(h)+4))
   info=wb.create_sheet('Report notes');info.append(['Tiny Tale',kind.title()]);info.append(['From (IST)',frm]);info.append(['To (IST)',to]);info.append(['Generated (IST)',stamp(datetime.utcnow())]);info.append(['Notes',note]);wb.save(buf)
   mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
  else:
   from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
   from reportlab.lib.pagesizes import A4, landscape
   from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
   from reportlab.lib import colors
   styles=getSampleStyleSheet();body=ParagraphStyle('cell',fontName='Helvetica',fontSize=8,leading=11,wordWrap='CJK');page=landscape(A4)
   doc=SimpleDocTemplate(buf,pagesize=page,rightMargin=28,leftMargin=28,topMargin=30,bottomMargin=32)
   story=[Paragraph('Tiny Tale | '+kind.title()+' report',styles['Title']),Paragraph(f'{frm} to {to} (IST) | Generated {stamp(datetime.utcnow())}',styles['Normal']),Spacer(1,12),Paragraph(escape(note),styles['Normal']),Spacer(1,16)]
   if data:
    table=Table([[Paragraph(escape(str(x))[:1500],body) for x in row] for row in [headers]+data],colWidths=[(page[0]-56)/len(headers)]*len(headers),repeatRows=1,splitInRow=1)
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dff0f6')),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f5f8fa')]),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#dce4e7')),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]));story.append(table)
   else:story.append(Paragraph('No records for this period.',styles['Normal']))
   def footer(canvas,doc):canvas.setFont('Helvetica',8);canvas.drawString(28,16,'Tiny Tale | Internal management report');canvas.drawRightString(page[0]-28,16,'Page '+str(doc.page))
   doc.build(story,onFirstPage=footer,onLaterPages=footer);mime='application/pdf'
  buf.seek(0);return send_file(buf,mimetype=mime,download_name=filename,as_attachment=request.args.get('download')=='1',max_age=0)
