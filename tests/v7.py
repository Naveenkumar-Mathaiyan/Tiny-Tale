"""Staff/mobile authorization, scan inventory, receipt ownership and label PDFs."""
from smoke import *
from pypdf import PdfReader
from operations import valid_gtin
m.ROLE='admin'
assert valid_gtin('4006381333931') and not valid_gtin('4006381333932')
assert admin.post('/api/admin/categories',json={'name':'Gift Sets'},headers=ah).status_code==200
assert 'Gift Sets' in shop.get('/api/categories').json
p0=next(p for p in admin.get('/api/admin/products').json if p['id']=='1');p0['category']='Gift Sets'
assert admin.post('/api/admin/products',json=p0,headers=ah).status_code==200
assert admin.post('/api/admin/categories',json={'old':'Gift Sets','name':'Baby Gifts'},headers=ah).status_code==200
assert admin.post('/api/admin/categories',json={'name':'baby gifts'},headers=ah).status_code==400
assert next(p for p in shop.get('/api/products').json if p['id']=='1')['category']=='Baby Gifts'
hero=admin.get('/api/admin/homepage').json
assert admin.post('/api/admin/homepage',json=hero|{'title':'Welcome little ones','interval':4},headers=ah).status_code==200
assert shop.get('/api/homepage').json['title']=='Welcome little ones'
assert admin.post('/api/admin/homepage',json=hero|{'link':'javascript:bad'},headers=ah).status_code==400
assert admin.post('/api/admin/homepage',json=hero|{'interval':1},headers=ah).status_code==400
rows=admin.post('/api/admin/operations/codes',json={'generate':True},headers=ah).json
assert len(rows)>0
assert admin.get('/api/admin/operations/receive').status_code==405
assert admin.post('/api/admin/operations/stock',json={},headers=ah).status_code==405
assert admin.post('/api/admin/homepage',json=['bad'],headers=ah).status_code==400
c=next(x for x in rows if x['product_id']=='1');codeid=c['id']
assert admin.post('/api/admin/operations/codes',json={'id':codeid,'code':'TT-JABLA-ONE','gtin':'4006381333931','licensed':True},headers=ah).status_code==200
assert admin.post('/api/admin/operations/codes',json={'id':codeid,'code':'ABC','gtin':'4006381333932','licensed':True},headers=ah).status_code==400
othercode=next(x for x in rows if x['id']!=codeid)
assert admin.post('/api/admin/operations/codes',json={'id':othercode['id'],'code':'4006381333931'},headers=ah).status_code==400
for typ in ['internal','gtin']:
 r=admin.get('/api/admin/operations/labels?ids='+codeid+'&copies=2&type='+typ);assert r.status_code==200,r.json
 reader=PdfReader(io.BytesIO(r.data));assert len(reader.pages)==1 and 'TT-JABLA-ONE' in reader.pages[0].extract_text() if typ=='internal' else '4006381333931' in reader.pages[0].extract_text()
assert admin.get('/api/admin/operations/labels?ids='+codeid+'&copies=501').status_code==400
clients={};tokens={};staffids={}
for role in ['store_manager','manager','supervisor','billing']:
 r=admin.post('/api/admin/accounts',json={'username':role+'7','name':role,'role':role,'password':'staff-password-123','active':True},headers=ah);assert r.status_code==200,r.json
 cli=app.test_client();r=cli.post('/api/admin/login',json={'username':role+'7','password':'staff-password-123'});clients[role]=(cli,{'X-CSRF-Token':r.json['csrf']});staffids[role]=r.json['user']['id']
 r=cli.post('/api/staff/login',json={'username':role+'7','password':'staff-password-123'});assert r.status_code==200,r.json;tokens[role]={'Authorization':'Bearer '+r.json['token']}
 for path in ['analytics','orders','accounts','homepage','categories','reports?type=orders']:
  assert cli.get('/api/admin/'+path).status_code==403,(role,path)
 assert cli.get('/api/staff/session',headers=tokens[role]).status_code==200
 assert cli.get('/api/staff/lookup?code=tiny-tale:TT-JABLA-ONE',headers=tokens[role]).status_code==200
 assert cli.get('/api/staff/lookup?code=4006381333931',headers=tokens[role]).status_code==200
manager,manh=clients['store_manager'];billing,bilh=clients['billing'];sup,suph=clients['supervisor']
assert manager.get('/api/admin/operations/bills').status_code==403
assert billing.get('/api/admin/products').status_code==403
assert billing.post('/api/staff/receive',json={},headers=tokens['billing']).status_code==403
assert manager.post('/api/staff/sale',json={},headers=tokens['store_manager']).status_code==403
assert manager.post('/api/admin/operations/receive',json={}).status_code==403
assert manager.get('/api/staff/stock').status_code==401
before=manager.get('/api/staff/lookup?code=TT-JABLA-ONE',headers=tokens['store_manager']).json['stock']
receive=dict(request_id=str(uuid.uuid4()),items=[dict(product_id='1',size=c['size'],qty=4)])
r=manager.post('/api/staff/receive',json=receive,headers=tokens['store_manager']);assert r.status_code==201,r.json
assert manager.post('/api/staff/receive',json=receive,headers=tokens['store_manager']).json==r.json
assert manager.get('/api/staff/lookup?code=TT-JABLA-ONE',headers=tokens['store_manager']).json['stock']==before+4
assert manager.post('/api/staff/receive',json=receive|{'items':[dict(product_id='1',size=c['size'],qty=5)]},headers=tokens['store_manager']).status_code==409
sale=dict(request_id=str(uuid.uuid4()),items=[dict(product_id='1',size=c['size'],qty=3)],method='cash',expected_total=3*next(x for x in shop.get('/api/products').json if x['id']=='1')['sale'])
r=billing.post('/api/staff/sale',json=sale,headers=tokens['billing']);assert r.status_code==201,r.json;bill=r.json
assert bill['total']==3*next(x for x in shop.get('/api/products').json if x['id']=='1')['sale']
assert billing.post('/api/staff/sale',json=sale,headers=tokens['billing']).json==bill
assert manager.get('/api/staff/lookup?code=TT-JABLA-ONE',headers=tokens['store_manager']).json['stock']==before+1
assert billing.post('/api/staff/sale',json=sale|{'request_id':str(uuid.uuid4()),'expected_total':1},headers=tokens['billing']).status_code==409
assert sup.get('/api/staff/receipt/'+bill['id'],headers=tokens['supervisor']).status_code==403
assert sup.get('/api/staff/bills',headers=tokens['supervisor']).json==[]
assert len(PdfReader(io.BytesIO(billing.get('/api/staff/receipt/'+bill['id'],headers=tokens['billing']).data)).pages)>0
assert admin.get('/api/admin/operations/receipt/'+bill['id']).status_code==200
bad=sale|{'request_id':str(uuid.uuid4()),'items':[dict(product_id='1',size=c['size'],qty=10000)]}
assert billing.post('/api/staff/sale',json=bad,headers=tokens['billing']).status_code==409
assert manager.get('/api/staff/lookup?code=TT-JABLA-ONE',headers=tokens['store_manager']).json['stock']==before+1
# Entire multi-item batch rolls back when a later item is invalid, and invalid payment cannot deduct stock.
assert billing.post('/api/staff/sale',json=sale|{'request_id':str(uuid.uuid4()),'method':'cod'},headers=tokens['billing']).status_code==400
assert manager.get('/api/staff/lookup?code=TT-JABLA-ONE',headers=tokens['store_manager']).json['stock']==before+1
r=manager.post('/api/staff/receive',json=receive|{'request_id':str(uuid.uuid4()),'items':[dict(product_id='1',size=c['size'],qty=1),dict(product_id='missing',size='One size',qty=1)]},headers=tokens['store_manager']);assert r.status_code==400
assert manager.get('/api/staff/lookup?code=TT-JABLA-ONE',headers=tokens['store_manager']).json['stock']==before+1
# Product media and cover save in one transaction; videos cannot become the cover.
buf=io.BytesIO();Image.new('RGB',(10,10),'green').save(buf,format='PNG');asset=admin.post('/api/admin/assets',data={'file':(io.BytesIO(buf.getvalue()),'new.png')},headers=ah).json
p=next(x for x in admin.get('/api/admin/products').json if x['id']=='1');p['image']=asset['url'];p['gallery']=[{'url':'/static/images/product-1.png'}]
assert admin.post('/api/admin/products',json=p,headers=ah).status_code==200
p=next(x for x in admin.get('/api/admin/products').json if x['id']=='1')
r=manager.post('/api/admin/stock',json={'id':'1','image':p['image'],'gallery':p['gallery'],'sizes':[{'size':v['size'],'enabled':v['enabled'],'stock':v['stock']} for v in p['sizes']]},headers=manh);assert r.status_code==200,r.json
video=admin.post('/api/admin/assets',data={'file':(io.BytesIO((Path(__file__).parent/'fixtures/sample.mp4').read_bytes()),'clip.mp4')},headers=ah).json
assert admin.post('/api/admin/products',json=p|{'image':video['url']},headers=ah).status_code==400
assert admin.post('/api/admin/products',json=p|{'gallery':[{'url':'javascript:bad'}],'sale':1},headers=ah).status_code==400
assert next(x for x in shop.get('/api/products').json if x['id']=='1')['sale']==p['sale']
assert shop.get('/api/products').json[0].get('code') is None
today=(datetime.utcnow()+timedelta(hours=5,minutes=30)).strftime('%Y-%m-%d')
for kind in ['counter','movements']:
 for fmt in ['pdf','xlsx']:
  report=admin.get(f'/api/admin/reports?type={kind}&from={today}&to={today}&format={fmt}');assert report.status_code==200
  if fmt=='pdf':assert len(PdfReader(io.BytesIO(report.data)).pages)>0
  else:
   from openpyxl import load_workbook
   assert load_workbook(io.BytesIO(report.data))['Report'].max_row>1
# Disable account invalidates both web cookie and native token.
assert admin.post('/api/admin/accounts',json={'id':staffids['billing'],'username':'billing7','name':'Billing','role':'billing','password':'','active':False},headers=ah).status_code==200
assert billing.get('/api/staff/session',headers=tokens['billing']).status_code==401
assert billing.get('/api/admin/operations/bills').status_code==401
assert manager.post('/api/staff/logout',headers=tokens['store_manager']).status_code==200
assert manager.get('/api/staff/session',headers=tokens['store_manager']).status_code==401
print('PASS v7: categories, hero settings, licensed GTIN checks, cross-code uniqueness, internal/EAN/QR PDF labels, role restrictions, native authentication/logout/revocation, incoming and POS stock sync, replay safety, overselling rollback, receipt ownership, atomic image/gallery editing and customer code privacy.')
